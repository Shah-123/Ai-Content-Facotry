"""Regression tests for the shared blog export service."""

from pathlib import Path

from exporters import export_all, export_to_docx, _iter_inline, _md_inline_to_rl


def test_export_all_writes_a_self_contained_html_file(tmp_path: Path):
    results = export_all(
        "# Export test\n\nA short paragraph.",
        tmp_path / "content" / "export_test",
        "Export Test",
        ["html"],
    )

    output_path = Path(results["html"])
    assert output_path.exists()
    assert "Export Test" in output_path.read_text(encoding="utf-8")


def test_inline_tokenizer_splits_bold_italic_code_and_links():
    spans = list(_iter_inline("plain **b** and *i* and `c` and [t](http://x)"))
    got = [(chunk, fmt) for chunk, fmt in spans if chunk.strip()]
    assert ("b", {"bold": True}) in got
    assert ("i", {"italic": True}) in got
    assert ("c", {"code": True}) in got
    assert ("t", {"url": "http://x"}) in got


def test_pdf_markup_preserves_formatting_and_escapes_text():
    out = _md_inline_to_rl("**bold** & <raw> [link](http://x)")
    assert "<b>bold</b>" in out
    assert '<a href="http://x"' in out
    assert "&amp;" in out and "&lt;raw&gt;" in out  # XML-escaped, not injected


def test_docx_bold_survives_as_a_run(tmp_path: Path):
    """The old renderer flattened **bold** to plain text; it must now be a bold run."""
    out = tmp_path / "content" / "fmt.docx"
    export_to_docx("Fmt", "This is **important** text.", out, job_dir=tmp_path)

    import docx
    doc = docx.Document(str(out))
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "important" in bold_runs


def _one_pixel_png() -> bytes:
    import base64
    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )


def test_every_format_embeds_the_same_pictures(tmp_path: Path):
    """DOCX and PDF used to hand-roll their own Markdown parser, so an escaped
    bracket in the alt text, an angle-bracket src or a <figure> block silently
    dropped the picture from those two formats while HTML kept it."""
    images = tmp_path / "assets" / "images"
    images.mkdir(parents=True)
    for name in ("a.png", "b.png", "c.png"):
        (images / name).write_bytes(_one_pixel_png())

    content = (
        "# Title\n\n"
        r"![Chart \[2024\] results](assets/images/a.png)" + "\n\n"
        "![Angle bracket src](<assets/images/b.png>)\n\n"
        '<figure><img src="assets/images/c.png" alt="Fig"><figcaption>Cap</figcaption></figure>\n'
    )

    from exporters import export_to_html, export_to_pdf
    assert export_to_html("Title", content, job_dir=tmp_path).count("data:image") == 3

    import docx
    docx_path = export_to_docx("Title", content, tmp_path / "out.docx", job_dir=tmp_path)
    assert len(docx.Document(str(docx_path)).inline_shapes) == 3

    # ReportLab raises on a missing file, so a clean build is the PDF's check.
    assert export_to_pdf("Title", content, tmp_path / "out.pdf", job_dir=tmp_path).exists()


def test_docx_renders_blocks_the_line_parser_used_to_mangle(tmp_path: Path):
    import docx

    content = (
        "# Title\n\n"
        "| Metric | Value |\n|---|---|\n| Speed | **90s** [src](http://x) |\n\n"
        "- bullet one\n- bullet two\n\n"
        "```\ncode block\n```\n\n"
        "> quoted line\n"
    )
    doc = docx.Document(str(export_to_docx("Title", content, tmp_path / "b.docx", job_dir=tmp_path)))

    table = doc.tables[0]
    assert (len(table.rows), len(table.columns)) == (2, 2)
    # Cells carry formatted runs now, not raw "**90s** [src](http://x)" markdown.
    assert "[src](http://x)" not in table.cell(1, 1).text
    assert "90s" in table.cell(1, 1).text

    texts = [p.text for p in doc.paragraphs]
    assert [p.style.name for p in doc.paragraphs if p.text == "bullet one"] == ["List Bullet"]
    assert "code block" in texts
    assert "quoted line" in texts
